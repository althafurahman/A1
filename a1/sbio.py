"""Dataset loading and answer-range plumbing.

Adapted from the official hackathon starter (ylookup/encode-hackathon, research/sb.py)
so the harness resolves tasks and answer cells exactly the way the shipped evaluator does.
"""

import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries


def load_dataset(dataset_dir):
    dataset_dir = Path(dataset_dir)
    tasks = json.loads((dataset_dir / "dataset.json").read_text())
    for t in tasks:
        t["id"] = str(t["id"])
        folder = dataset_dir / t["spreadsheet_path"]
        t["init_xlsx"] = str(next(folder.glob("*init*.xlsx")))
        golden = next(folder.glob("*golden*.xlsx"), None)
        t["golden_xlsx"] = str(golden) if golden else None
    return tasks


def parse_answer_position(answer_position):
    cleaned = answer_position.replace("'", "").replace('"', "")
    tokens = [cleaned] if cleaned.count("!") == 1 else cleaned.split(",")
    parsed = []
    for token in tokens:
        token = token.strip()
        if "!" in token:
            sheet, rng = token.rsplit("!", 1)
            parsed.append((sheet, _repair_range(rng)))
        else:
            parsed.append((None, _repair_range(token)))
    return parsed


def _repair_range(rng):
    if ":" not in rng:
        return rng
    start, end = rng.split(":", 1)
    if end.isdigit():
        col = "".join(ch for ch in start if ch.isalpha())
        return f"{start}:{col}{end}"
    return rng


def expand_range(cell_range, max_row=None):
    """Expand A1:B3 to cell coordinates. Whole-column ranges like A:G need max_row."""
    min_col, min_row, max_col, last_row = range_boundaries(cell_range)
    min_row = min_row or 1
    last_row = last_row or max_row or min_row
    return [f"{get_column_letter(c)}{r}" for r in range(min_row, last_row + 1) for c in range(min_col, max_col + 1)]


def answer_ranges(task):
    return [(sheet or task.get("answer_sheet"), rng) for sheet, rng in parse_answer_position(task["answer_position"])]


def answer_cells(task, wb=None):
    cells = []
    for sheet, rng in answer_ranges(task):
        max_row = None
        if wb is not None:
            ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
            max_row = ws.max_row
        cells.extend((sheet, coord) for coord in expand_range(rng, max_row))
    return cells


def read_answer_region(path, task, max_cells=600):
    """Values currently in the answer region of a workbook, as printable lines."""
    wb = openpyxl.load_workbook(path)  # not data_only: shows formulas the model wrongly wrote
    lines, n = [], 0
    for sheet, coord in answer_cells(task, wb):
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        v = ws[coord].value
        if v is not None:
            lines.append(f"{ws.title}!{coord} = {v!r}")
        n += 1
        if len(lines) >= max_cells:
            lines.append("... (truncated)")
            break
    if not lines:
        lines.append(f"(all {n} cells in the answer region are empty)")
    return "\n".join(lines)


def read_jsonl(path):
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]
