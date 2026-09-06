"""Per-task pipeline: code agent with repair loop and self-verification, direct-answer fallback."""

import json
import re
import shutil
import time
from pathlib import Path


def clean_status(text: str, cap: int = 180) -> str:
    return " ".join(str(text).split())[:cap]

import openpyxl

from .coderun import extract_code, run_code
from .tinker_llm import LLMError
from .sbio import answer_cells, read_answer_region
from .serialize import serialize_workbook

VALUE_RULES = """Grading rules for the answer cells (only these cells are compared, by VALUE, after recalculation):
- Write plain computed VALUES, never formula strings. Numbers are compared rounded to 2 decimals.
- Dates must be real date/datetime objects (or the exact Excel serial number), never text like "2024-03-01".
- Booleans must be real booleans (True/False), not 1/0 and not "TRUE".
- A cell that should be empty must be empty (None). Empty string equals empty.
- Text must match exactly (case, spacing, punctuation), except text that parses as a number is compared numerically.
- Never write leading or trailing whitespace in a text value; strip() strings you assemble or copy.
- Do not change the sheet names of the workbook."""

CODE_SYSTEM = """You are an expert spreadsheet engineer. You receive a serialized Excel workbook and a user instruction from an Excel forum. Instead of answering in text, you write ONE Python script that computes the result from the real workbook and writes it into the answer region.

Environment: Python 3.12 with openpyxl, pandas and numpy. Two names are predefined for you: INIT (path of the input workbook) and OUT (path your script must create). No network access.

Script template you should follow:
import shutil, openpyxl
shutil.copy(INIT, OUT)
vals = openpyxl.load_workbook(INIT, data_only=True)   # read cached VALUES here
wb = openpyxl.load_workbook(OUT)                       # write here (keeps formulas elsewhere intact)
ws = wb["<answer sheet>"]
# ... compute, then assign plain values into the answer cells ...
wb.save(OUT)

Rules:
- {value_rules}
- Read cell values from the data_only load (or pandas.read_excel). Formula cells there hold their cached results.
- The instruction may ask for a formula, VBA or a manual technique. Ignore the requested mechanism: produce the final VALUES that would result in the answer region.
- Compute from the actual data. Never hardcode results you read off the serialization for large ranges; loop over the real cells.
- Only the answer region is graded; everything else in the workbook is ignored. But do not delete or rename sheets.
- If the answer sheet does not exist in the workbook, create it with exactly that name and write there.
- Print short diagnostics (row counts, a few computed values) so failures can be debugged.
- Be deterministic. End by saving OUT.

Reply with a short plan (3 sentences max) followed by exactly one ```python code block."""

REPAIR_USER = """Your script failed. Fix it and reply with one ```python code block (full script, not a diff).

stderr:
{stderr}

stdout:
{stdout}"""

VERIFY_SYSTEM = """You are a meticulous spreadsheet QA reviewer. You get a user instruction, workbook context, and the values a candidate solution wrote into the graded answer region. Decide whether the values plausibly satisfy the instruction.

Check: correct interpretation of the instruction, sane magnitudes, right data types (dates as dates, numbers as numbers), no leftover formula strings, region actually filled where it should be, empty where it should be empty. If the dump lists EMPTY cells, decide whether the instruction really implies those exact cells stay empty — partially filled regions are the most common near-miss.

Reply with JSON only: {"pass": true} or {"pass": false, "issue": "<one concrete sentence on what is wrong and how to fix it>"}"""

DIRECT_SYSTEM = """You are a spreadsheet expert. You get a serialized workbook and a user instruction. Compute the final values the answer range must contain after the instruction is applied.

{value_rules}

Reply with JSON only, no prose: {{"cells": [{{"cell": "B6", "value": 42}}, {{"cell": "B7", "value": null}}]}}
One entry per cell in the answer range. Use null for cells that must be empty. Dates as "DATE(2024,3,1)" strings are not allowed — write the Excel serial number instead."""

MAX_REPAIRS = 2


def build_task_user(task, serialization):
    sheet = task.get("answer_sheet") or "(the active sheet)"
    return (
        f"## Instruction\n{task['instruction']}\n\n"
        f"## Workbook\n{serialization}\n\n"
        f"## Answer region\nSheet: {sheet}\nCells: {task['answer_position']}\n"
        f"Data region hint: {task.get('data_position') or 'n/a'}\n"
    )


def parse_json_reply(text: str) -> dict:
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S)
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < 0:
        raise ValueError(f"no JSON in reply: {text[:120]!r}")
    return json.loads(text[start:end + 1])


def write_direct_answer(task, cells: list[dict], out_path: Path):
    values = {str(c.get("cell", "")).upper(): c.get("value") for c in cells}
    shutil.copy(task["init_xlsx"], out_path)
    wb = openpyxl.load_workbook(out_path)
    for sheet, coord in answer_cells(task, wb):
        if sheet and sheet not in wb.sheetnames:
            wb.create_sheet(sheet)  # grading reads this sheet by name; never write elsewhere
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        if coord in values:
            ws[coord] = values[coord]
    wb.save(out_path)


def hygiene_issues(out_path: Path, task) -> str | None:
    """Mechanical checks on the written answer region: formula strings, stray whitespace."""
    wb = openpyxl.load_workbook(out_path)
    formulas, padded = [], []
    for sheet, coord in answer_cells(task, wb):
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        v = ws[coord].value
        if isinstance(v, str):
            if v.startswith("="):
                formulas.append(f"{ws.title}!{coord}")
            elif v != v.strip() and v.strip():
                padded.append(f"{ws.title}!{coord}")
    problems = []
    if formulas:
        problems.append(f"You wrote formula STRINGS into answer cells ({', '.join(formulas[:5])}...). "
                        "Write computed plain values instead.")
    if padded:
        problems.append(f"Text values with leading/trailing whitespace in {', '.join(padded[:8])}"
                        f"{'...' if len(padded) > 8 else ''}. Strip whitespace; grading compares text exactly.")
    return " ".join(problems) or None


class TaskRunner:
    """Runs one task, accumulating trace lines. Not thread-safe; one instance per task."""

    def __init__(self, client, task, out_dir: Path, strategy: str = "full"):
        self.client = client
        self.task = task
        self.out_dir = out_dir
        self.strategy = strategy
        self.out_xlsx = out_dir / "outputs" / f"{task['id']}.xlsx"
        self.traces = []
        self.step = 0

    def _trace_llm(self, record):
        self.step += 1
        self.traces.append({"step": self.step, **record})

    def _trace_tool(self, name, tool_input, tool_output):
        self.step += 1
        self.traces.append({
            "step": self.step, "model": None, "prompt": None, "response": None,
            "input_tokens": None, "output_tokens": None, "latency_ms": None, "error": None,
            "tool": name, "tool_input": tool_input[:20_000], "tool_output": tool_output[:8_000],
        })

    async def _call(self, system, user, phase, max_tokens=24_576):
        try:
            rec = await self.client.complete(system, user, max_tokens=max_tokens)
            rec["phase"] = phase
            self._trace_llm(rec)
            return rec["response"]
        except LLMError as e:
            self.step += 1
            self.traces.append({"step": self.step, "model": self.client.model, "prompt": (system + "\n\n" + user)[:20_000],
                               "response": None, "input_tokens": None, "output_tokens": None,
                               "latency_ms": None, "error": str(e)[:500], "phase": phase})
            raise

    SERIALIZATION_BUDGETS = (120_000, 45_000, 16_000)  # chars; shrink when the prompt overflows the context

    @staticmethod
    def _prompt_overflow(status: str | None) -> bool:
        return bool(status) and ("prompt too long" in status or "context window" in status)

    async def solve(self) -> str:
        task = self.task
        started = time.time()
        wb_vals = openpyxl.load_workbook(task["init_xlsx"], data_only=True)
        wb_form = openpyxl.load_workbook(task["init_xlsx"])

        status = None
        for budget in self.SERIALIZATION_BUDGETS:
            serialization = serialize_workbook(wb_vals, wb_form, task, char_budget=budget)
            task_user = build_task_user(task, serialization)
            if self.strategy in ("full", "code"):
                status = await self._solve_code(task_user)
            if status != "ok" and self.strategy in ("full", "direct") and not self._prompt_overflow(status):
                status = await self._solve_direct(task_user, note=status)
            if not self._prompt_overflow(status):
                break
        if status != "ok":
            shutil.copy(task["init_xlsx"], self.out_xlsx)
            status = status or "error: no strategy produced output"
        self._trace_tool("done", f"strategy={self.strategy}", f"status={status} elapsed={int(time.time()-started)}s")
        return status

    async def _solve_code(self, task_user) -> str:
        system = CODE_SYSTEM.format(value_rules=VALUE_RULES)
        try:
            reply = await self._call(system, task_user, "code")
        except LLMError as e:
            return f"error: {clean_status(e)}"
        code = extract_code(reply)
        produced_any = False
        token_budget = 24_576
        for attempt in range(MAX_REPAIRS + 1):
            if code is None:
                # usually mid-think truncation on hard tasks: give the retry more room to finish
                token_budget = 45_000
                result = {"ok": False, "stdout": "",
                          "stderr": "Your reply contained no ```python code block. "
                                    "Reply with the complete script in exactly one ```python block."}
            else:
                result = run_code(code, self.task["init_xlsx"], str(self.out_xlsx))
                self._trace_tool("python", code, f"ok={result['ok']}\nstdout:\n{result['stdout']}\nstderr:\n{result['stderr']}")
            if result["ok"]:
                produced_any = True
                issue = hygiene_issues(self.out_xlsx, self.task)
                if issue is None:
                    verified = await self._verify(task_user)
                    if verified is True:
                        return "ok"
                    issue = f"Output was produced, but review found a problem: {verified}"
                result = {"ok": False, "stdout": result["stdout"], "stderr": issue}
            if attempt == MAX_REPAIRS:
                if produced_any:
                    return "ok"  # advisory issues at this point; a produced output beats the fallback
                return f"error: code failed: {clean_status(result['stderr'], 150)}"
            try:
                reply = await self._call(system, task_user + "\n\n" + REPAIR_USER.format(**result), "repair",
                                         max_tokens=token_budget)
            except LLMError as e:
                return "ok" if produced_any else f"error: {clean_status(e)}"
            code = extract_code(reply)
        return "error: unreachable"

    async def _verify(self, task_user):
        """True if the reviewer passes the output; otherwise the issue text."""
        region = read_answer_region(self.out_xlsx, self.task)
        before = read_answer_region(self.task["init_xlsx"], self.task)
        user = (f"{task_user}\n\n## Answer region BEFORE (in the original workbook)\n{before}\n\n"
                f"## Answer region AFTER (what the candidate wrote)\n{region}\n\n"
                f"Compare before and after against the instruction. Watch for content shifted by a row "
                f"or column relative to where the instruction and the original layout say it belongs. "
                f"Does the AFTER state satisfy the instruction?")
        try:
            reply = await self._call(VERIFY_SYSTEM, user, "verify")
            verdict = parse_json_reply(reply)
            if verdict.get("pass") is True:
                return True
            return str(verdict.get("issue") or "reviewer failed the output without a reason")
        except (LLMError, ValueError, json.JSONDecodeError):
            return True  # a broken reviewer must not sink a produced answer

    async def _solve_direct(self, task_user, note=None) -> str:
        system = DIRECT_SYSTEM.format(value_rules=VALUE_RULES)
        try:
            reply = await self._call(system, task_user, "direct")
            answer = parse_json_reply(reply)
            write_direct_answer(self.task, answer.get("cells", []), self.out_xlsx)
            return "ok"
        except (LLMError, ValueError, json.JSONDecodeError, KeyError, TypeError) as e:
            prior = f"{note}; " if note else ""
            return clean_status(f"error: {prior}direct fallback failed: {e}", 250)
