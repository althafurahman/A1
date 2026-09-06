"""Build a supervised JSONL from a harness run: prompt (rebuilt exactly as inference) -> plan + working script.
    uv run --project research python scripts/build_sft.py experiments/full-400-001 experiments/sft-001
Keeps only tasks that PASS in <run>/results.json and are not in experiments/heldout78.ids.
Target = the model reply (phase code/repair) that produced the last successful python tool step.
Filters: script must read the workbook (no hard-coded answers); prompt+target <= MAX_TOKENS.
"""
import json, random, sys, os
from pathlib import Path
sys.path.insert(0, ".")
import openpyxl
from a1.harness import CODE_SYSTEM, VALUE_RULES, build_task_user
from a1.serialize import serialize_workbook
from a1.sbio import load_dataset

run, out = Path(sys.argv[1]), Path(sys.argv[2])
dataset = sys.argv[3] if len(sys.argv) > 3 else "research/data/spreadsheetbench_verified_400"
MAX_TOKENS = 15_000
heldout = set(open("experiments/heldout78.ids").read().strip().split(","))
passed = {i["id"] for i in json.load(open(run / "results.json"))["items"] if i.get("pass")}
tasks = {t["id"]: t for t in load_dataset(dataset)}
from tinker_cookbook.tokenizer_utils import get_tokenizer
tok = get_tokenizer("Qwen/Qwen3.8-27B")
system = CODE_SYSTEM.format(value_rules=VALUE_RULES)

def target_reply(recs):
    """Reply that produced the last ok python step."""
    last = None
    for i, r in enumerate(recs):
        if r.get("tool") == "python" and r.get("tool_output", "").startswith("ok=True"):
            last = i
    if last is None: return None
    for r in reversed(recs[:last]):
        if r.get("model") and r.get("phase") in ("code", "repair") and r.get("response"):
            return r["response"]
    return None

out.mkdir(parents=True, exist_ok=True)
kept, skipped = [], {"heldout": 0, "not_pass": 0, "no_reply": 0, "no_read": 0, "too_long": 0}
for f in sorted((run / "traces").glob("*.jsonl")):
    tid = f.stem
    if tid in heldout: skipped["heldout"] += 1; continue
    if tid not in passed: skipped["not_pass"] += 1; continue
    recs = [json.loads(l) for l in open(f) if l.strip()]
    reply = target_reply(recs)
    if not reply: skipped["no_reply"] += 1; continue
    if "load_workbook" not in reply and "read_excel" not in reply: skipped["no_read"] += 1; continue
    t = tasks[tid]
    user = build_task_user(t, serialize_workbook(openpyxl.load_workbook(t["init_xlsx"], data_only=True),
                                                 openpyxl.load_workbook(t["init_xlsx"]), t))
    n = len(tok.encode(system + user + reply))
    if n > MAX_TOKENS: skipped["too_long"] += 1; continue
    kept.append({"id": tid, "tokens": n, "messages": [{"role": "system", "content": system},
                                                       {"role": "user", "content": user},
                                                       {"role": "assistant", "content": reply}]})
random.seed(0); random.shuffle(kept)
with open(out / "train.jsonl", "w") as fh:
    for ex in kept: fh.write(json.dumps({"messages": ex["messages"]}, ensure_ascii=False) + "\n")
open(out / "train_ids.txt", "w").write(",".join(ex["id"] for ex in kept) + "\n")
toks = sorted(ex["tokens"] for ex in kept)
print(f"kept {len(kept)}  skipped {skipped}")
if toks: print(f"tokens/example min {toks[0]} median {toks[len(toks)//2]} max {toks[-1]} total {sum(toks)}")
print("\n=== 10 samples for hand audit (id, plan line, first 12 code lines) ===")
for ex in kept[:10]:
    reply = ex["messages"][2]["content"]; code = reply.split("```")[1] if "```" in reply else reply
    print(f"\n--- {ex['id']} ({ex['tokens']} tok) ---\n{reply.split('```')[0].strip()[:300]}\n" + "\n".join(code.splitlines()[1:13]))
